"""
export_dialog.py — ExportDialog: Unified dialog for DOCX and BibTeX export.
"""
import re
import unicodedata
from pathlib import Path

from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QFormLayout,
    QGroupBox, QLineEdit, QComboBox, QSpinBox, QCheckBox,
    QTextEdit, QLabel, QPushButton, QFileDialog, QMessageBox,
)

from app.theme import icon
from app.citation_engine import CitationStyleEngine
from app.adapter import EntryAdapter
from app.i18n import LanguageManager
from app.dialogs.style_editor import CustomStyleEditorDialog


class ExportDialog(QDialog):
    """Configure and generate a bibliography export file (DOCX or BibTeX)."""

    _TYPE_MAP = {
        "article":          "@article",
        "book":             "@book",
        "chapter":          "@incollection",
        "triennale":        "@mastersthesis",
        "magistrale":       "@mastersthesis",
        "dottorato":        "@phdthesis",
        "specializzazione": "@mastersthesis",
        "tesi":             "@mastersthesis",
    }

    def __init__(self, adapter: EntryAdapter, lang_mgr: LanguageManager, parent=None):
        super().__init__(parent)
        self._adapter = adapter
        self._lang = lang_mgr
        self.setWindowTitle(self._lang.tr("tb_export"))
        self.setWindowIcon(icon("file-export"))
        self.setMinimumSize(640, 600)
        self._build_ui()
        self._load_presets()

    def _load_presets(self):
        from PySide6.QtCore import QSettings
        s = QSettings("CiteMind", "CiteMind")
        fmt_idx = s.value("export_format", 0, type=int)
        self._format_combo.setCurrentIndex(fmt_idx if fmt_idx < self._format_combo.count() else 0)
        
        type_idx = s.value("export_type_filter", 0, type=int)
        self._type_combo.setCurrentIndex(type_idx if type_idx < self._type_combo.count() else 0)
        
        self._abbrev_check.setChecked(s.value("export_abbrev", "false") == "true")
        
        style = s.value("export_style", "")
        if style:
            self._style_combo.setCurrentText(style)
            
        font = s.value("export_font", "Times New Roman")
        self._font_combo.setCurrentText(font)
        
        size = s.value("export_size", 12, type=int)
        self._size_spin.setValue(size)
        
        self._sorted_check.setChecked(s.value("export_sorted", "true") == "true")

    def _save_presets(self):
        from PySide6.QtCore import QSettings
        s = QSettings("CiteMind", "CiteMind")
        s.setValue("export_format", self._format_combo.currentIndex())
        s.setValue("export_type_filter", self._type_combo.currentIndex())
        s.setValue("export_abbrev", "true" if self._abbrev_check.isChecked() else "false")
        s.setValue("export_style", self._style_combo.currentText())
        s.setValue("export_font", self._font_combo.currentText())
        s.setValue("export_size", self._size_spin.value())
        s.setValue("export_sorted", "true" if self._sorted_check.isChecked() else "false")

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setSpacing(12)
        root.setContentsMargins(20, 20, 20, 20)

        grp = QGroupBox(self._lang.tr("docx_export_group"))
        form = QFormLayout(grp)
        form.setSpacing(8)

        # Common
        self._format_combo = QComboBox()
        self._format_combo.addItems(["DOCX Document", "BibTeX File", "BibLaTeX File", "PDF Report"])

        self._type_combo = QComboBox()
        self._type_combo.addItems([
            self._lang.tr("docx_export_filter_all"),
            self._lang.tr("docx_export_filter_bib"),
            self._lang.tr("docx_export_filter_sit"),
        ])

        self._abbrev_check = QCheckBox(self._lang.tr("docx_export_abbrev"))
        self._abbrev_check.setChecked(False)

        # DOCX Specific
        self._title_edit = QLineEdit("References")
        self._style_combo = QComboBox()
        self._style_combo.addItems(CitationStyleEngine.get_all_styles())

        self._btn_edit_styles = QPushButton(icon("settings"), "")
        self._btn_edit_styles.setToolTip(self._lang.tr("docx_export_style_tip"))
        self._btn_edit_styles.clicked.connect(self._open_style_editor)

        style_row = QHBoxLayout()
        style_row.setContentsMargins(0, 0, 0, 0)
        style_row.addWidget(self._style_combo, 1)
        style_row.addWidget(self._btn_edit_styles)

        self._font_combo = QComboBox()
        self._font_combo.addItems(["Times New Roman", "Calibri", "Arial", "Georgia", "Garamond"])

        self._size_spin = QSpinBox()
        self._size_spin.setRange(8, 20)
        self._size_spin.setValue(12)

        self._sorted_check = QCheckBox(self._lang.tr("docx_export_sort"))
        self._sorted_check.setChecked(True)

        form.addRow("Formato di Esportazione", self._format_combo)
        form.addRow(self._lang.tr("docx_export_include"), self._type_combo)
        form.addRow(self._lang.tr("docx_export_doc_title"), self._title_edit)
        form.addRow(self._lang.tr("docx_export_style"), style_row)
        form.addRow(self._lang.tr("docx_export_font"), self._font_combo)
        form.addRow(self._lang.tr("docx_export_size"), self._size_spin)
        form.addRow(self._sorted_check)
        form.addRow(self._abbrev_check)
        root.addWidget(grp)

        root.addWidget(QLabel(self._lang.tr("docx_export_preview")))
        self._preview = QTextEdit()
        self._preview.setReadOnly(True)
        self._preview.setMaximumHeight(250)
        root.addWidget(self._preview)

        btn_row = QHBoxLayout()
        prev_btn = QPushButton(icon("eye"), self._lang.tr("docx_export_btn_refresh"))
        export_btn = QPushButton(icon("file-export"), self._lang.tr("docx_export_btn_export"))
        export_btn.setObjectName("btnPrimary")
        cancel_btn = QPushButton(self._lang.tr("docx_export_btn_cancel"))
        btn_row.addWidget(prev_btn)
        btn_row.addStretch()
        btn_row.addWidget(cancel_btn)
        btn_row.addWidget(export_btn)
        root.addLayout(btn_row)

        # Signals
        prev_btn.clicked.connect(self._update_preview)
        export_btn.clicked.connect(self._do_export)
        cancel_btn.clicked.connect(self.reject)

        self._format_combo.currentIndexChanged.connect(self._on_format_changed)
        self._style_combo.currentIndexChanged.connect(self._update_preview)
        self._type_combo.currentIndexChanged.connect(self._update_preview)
        self._abbrev_check.stateChanged.connect(self._update_preview)

        self._on_format_changed()

    def _on_format_changed(self):
        fmt_idx = self._format_combo.currentIndex()
        is_docx_or_pdf = fmt_idx == 0 or fmt_idx == 3
        self._title_edit.setVisible(is_docx_or_pdf)
        self._style_combo.setVisible(is_docx_or_pdf)
        self._btn_edit_styles.setVisible(is_docx_or_pdf)
        self._font_combo.setVisible(is_docx_or_pdf)
        self._size_spin.setVisible(is_docx_or_pdf)
        self._sorted_check.setVisible(is_docx_or_pdf)
        
        # Hide their labels as well
        form = self._format_combo.parent().layout()
        
        # In QFormLayout, labelForField works on the field layout widget
        # Let's do it safely
        for i in range(form.rowCount()):
            item = form.itemAt(i, QFormLayout.FieldRole)
            if item:
                w = item.widget() or item.layout()
                label = form.itemAt(i, QFormLayout.LabelRole)
                if label and label.widget():
                    if w == self._title_edit:
                        label.widget().setVisible(is_docx_or_pdf)
                    elif hasattr(w, 'indexOf') and w.indexOf(self._style_combo) != -1:
                        label.widget().setVisible(is_docx_or_pdf)
                    elif w == self._font_combo:
                        label.widget().setVisible(is_docx_or_pdf)
                    elif w == self._size_spin:
                        label.widget().setVisible(is_docx_or_pdf)

        if is_docx_or_pdf:
            self._preview.setStyleSheet("")
        else:
            self._preview.setStyleSheet("font-family: 'Consolas', 'Courier New', monospace; font-size: 12px;")

        self._update_preview()

    # ── helpers ───────────────────────────────────────────────────────────
    def _abbreviate_authors(self, authors_str: str) -> str:
        if not authors_str:
            return ""
        abbreviated = []
        for author in [a.strip() for a in authors_str.split(";") if a.strip()]:
            parts = author.split(",", 1)
            if len(parts) == 2:
                last       = parts[0].strip()
                first_abbr = []
                for fn in parts[1].strip().split():
                    if "-" in fn:
                        first_abbr.append("-".join(
                            sfn[0].upper() + "." for sfn in fn.split("-") if sfn
                        ))
                    else:
                        clean = fn.strip(".")
                        if clean:
                            first_abbr.append(clean[0].upper() + ".")
                abbreviated.append(f"{last}, {' '.join(first_abbr)}")
            else:
                abbreviated.append(author)
        return "; ".join(abbreviated)

    def _get_entries(self) -> list[dict]:
        idx = self._type_combo.currentIndex()
        et  = "" if idx == 0 else ("bibliography" if idx == 1 else "sitography")
        return self._adapter.get_all_for_export(et)

    # DOCX Helpers
    def _current_style(self) -> str:
        return self._style_combo.currentText()

    def _open_style_editor(self):
        dlg  = CustomStyleEditorDialog(self._lang, self)
        dlg.exec()
        curr = self._style_combo.currentText()
        self._style_combo.clear()
        self._style_combo.addItems(CitationStyleEngine.get_all_styles())
        if curr in CitationStyleEngine.get_all_styles():
            self._style_combo.setCurrentText(curr)

    # BibTeX Helpers
    @staticmethod
    def _strip_accents(s: str) -> str:
        nfkd = unicodedata.normalize("NFKD", s)
        return "".join(c for c in nfkd if unicodedata.category(c) != "Mn")

    @classmethod
    def _make_cite_key(cls, entry: dict) -> str:
        authors = entry.get("authors", "").strip()
        year = entry.get("year", "").strip()
        clean_year = re.sub(r"[^0-9a-zA-Z]", "", year)

        if authors:
            first_author = authors.split(";")[0].strip()
            surname = first_author.split(",")[0].strip()
            surname = cls._strip_accents(surname)
            surname = re.sub(r"[^a-zA-Z]", "", surname)
        else:
            title = entry.get("title", "untitled").strip()
            surname = cls._strip_accents(title.split()[0]) if title else "untitled"
            surname = re.sub(r"[^a-zA-Z]", "", surname)

        return f"{surname}{clean_year}"

    @classmethod
    def _deduplicate_keys(cls, entries: list[dict]) -> list[str]:
        base_keys = [cls._make_cite_key(e) for e in entries]
        counts: dict[str, int] = {}
        for k in base_keys:
            counts[k] = counts.get(k, 0) + 1

        used: dict[str, int] = {}
        result = []
        for k in base_keys:
            if counts[k] > 1:
                idx = used.get(k, 0)
                suffix = chr(ord("a") + idx)
                used[k] = idx + 1
                result.append(f"{k}{suffix}")
            else:
                result.append(k)
        return result

    @classmethod
    def _bibtex_type(cls, entry: dict) -> str:
        entry_type = entry.get("entry_type", "")
        pub_type = entry.get("pub_type", "").strip()

        if entry_type == "sitography":
            return "@online"
        if entry_type == "tesi" or pub_type in (
            "triennale", "magistrale", "dottorato", "specializzazione", "tesi"
        ):
            return cls._TYPE_MAP.get(pub_type, "@mastersthesis")

        if pub_type in cls._TYPE_MAP:
            return cls._TYPE_MAP[pub_type]

        journal = entry.get("journal", "").strip()
        publisher = entry.get("publisher", "").strip()
        if journal and not publisher:
            return "@article"
        if publisher and not journal:
            return "@book"
        if journal:
            return "@article"
        return "@misc"

    _BIBLATEX_TYPE_MAP = {
        "article":          "@article",
        "book":             "@book",
        "chapter":          "@incollection",
        "triennale":        "@thesis",
        "magistrale":       "@thesis",
        "dottorato":        "@thesis",
        "specializzazione": "@thesis",
        "tesi":             "@thesis",
    }

    _BIBLATEX_THESIS_TYPE = {
        "triennale":        "bathesis",
        "magistrale":       "mathesis",
        "dottorato":        "phdthesis",
        "specializzazione": "mathesis",
        "tesi":             "mathesis",
    }

    @classmethod
    def _biblatex_type(cls, entry: dict) -> tuple[str, str]:
        """Return (biblatex_entry_type, thesis_type_field_or_empty)."""
        entry_type = entry.get("entry_type", "")
        pub_type = entry.get("pub_type", "").strip()

        if entry_type == "sitography":
            return "@online", ""

        if entry_type == "tesi" or pub_type in cls._BIBLATEX_THESIS_TYPE:
            thesis_type = cls._BIBLATEX_THESIS_TYPE.get(pub_type, "mathesis")
            return "@thesis", thesis_type

        if pub_type in cls._BIBLATEX_TYPE_MAP:
            return cls._BIBLATEX_TYPE_MAP[pub_type], ""

        journal = entry.get("journal", "").strip()
        publisher = entry.get("publisher", "").strip()
        if journal and not publisher:
            return "@article", ""
        if publisher and not journal:
            return "@book", ""
        if journal:
            return "@article", ""
        return "@misc", ""

    @classmethod
    def _authors_to_bibtex(cls, authors_str: str) -> str:
        if not authors_str:
            return ""
        parts = [a.strip() for a in authors_str.split(";") if a.strip()]
        return " and ".join(parts)

    @classmethod
    def _escape_bibtex(cls, value: str) -> str:
        return value.replace("&", r"\&").replace("%", r"\%").replace("#", r"\#")

    def _format_entry_bibtex(self, entry: dict, cite_key: str, abbrev: bool = False) -> str:
        btype = self._bibtex_type(entry)
        lines = [f"{btype}{{{cite_key},"]

        authors = entry.get("authors", "").strip()
        if abbrev and authors:
            authors = self._abbreviate_authors(authors)

        field_map = [
            ("author",    self._authors_to_bibtex(authors)),
            ("title",     "{" + self._escape_bibtex(entry.get("title", "")) + "}"),
            ("year",      entry.get("year", "").strip()),
            ("journal",   entry.get("journal", "").strip()),
            ("publisher", entry.get("publisher", "").strip()),
            ("volume",    entry.get("volume", "").strip()),
            ("number",    entry.get("issue", "").strip()),
            ("pages",     entry.get("pages", "").strip().replace("-", "--")),
            ("edition",   entry.get("edition", "").strip()),
            ("doi",       entry.get("doi", "").strip()),
            ("url",       entry.get("url", "").strip()),
            ("urldate",   entry.get("access_date", "").strip()),
            ("note",      entry.get("notes", "").strip()),
            ("address",   entry.get("location", "").strip()),
        ]

        if btype in ("@mastersthesis", "@phdthesis"):
            pub = entry.get("publisher", "").strip()
            if pub:
                field_map = [(k, v) if k != "publisher" else ("school", pub) for k, v in field_map]

        for field_name, value in field_map:
            if not value or value == "{}":
                continue
            if field_name == "title":
                lines.append(f"  {field_name} = {value},")
            else:
                lines.append(f"  {field_name} = {{{value}}},")

        lines.append("}")
        return "\n".join(lines)

    def _format_entry_biblatex(self, entry: dict, cite_key: str, abbrev: bool = False) -> str:
        """Format a single entry in BibLaTeX format.

        Uses BibLaTeX-specific entry types (@thesis with type field,
        @online for web) and field names (date, journaltitle, location,
        institution) for full compatibility with the biblatex package.
        """
        btype, thesis_type = self._biblatex_type(entry)
        lines = [f"{btype}{{{cite_key},"]

        authors = entry.get("authors", "").strip()
        if abbrev and authors:
            authors = self._abbreviate_authors(authors)

        # BibLaTeX uses 'date' instead of 'year', 'journaltitle' instead of 'journal',
        # 'location' instead of 'address', 'institution' for theses.
        field_map = [
            ("author",       self._authors_to_bibtex(authors)),
            ("title",        "{" + self._escape_bibtex(entry.get("title", "")) + "}"),
            ("date",         entry.get("year", "").strip()),
            ("journaltitle", entry.get("journal", "").strip()),
            ("publisher",    entry.get("publisher", "").strip()),
            ("volume",       entry.get("volume", "").strip()),
            ("number",       entry.get("issue", "").strip()),
            ("pages",        entry.get("pages", "").strip().replace("-", "--")),
            ("edition",      entry.get("edition", "").strip()),
            ("doi",          entry.get("doi", "").strip()),
            ("url",          entry.get("url", "").strip()),
            ("urldate",      entry.get("access_date", "").strip()),
            ("note",         entry.get("notes", "").strip()),
            ("location",     entry.get("location", "").strip()),
        ]

        # For thesis entries, use 'institution' instead of 'publisher'
        if btype == "@thesis":
            pub = entry.get("publisher", "").strip()
            if pub:
                field_map = [
                    (k, v) if k != "publisher" else ("institution", pub)
                    for k, v in field_map
                ]
            if thesis_type:
                field_map.append(("type", thesis_type))

        # For @incollection, add booktitle
        if btype == "@incollection":
            book_title = entry.get("book_title", "").strip()
            if book_title:
                field_map.append(("booktitle", "{" + self._escape_bibtex(book_title) + "}"))
            editors = entry.get("editors", "").strip()
            if editors:
                field_map.append(("editor", self._authors_to_bibtex(editors)))

        for field_name, value in field_map:
            if not value or value == "{}":
                continue
            if field_name in ("title", "booktitle"):
                lines.append(f"  {field_name} = {value},")
            else:
                lines.append(f"  {field_name} = {{{value}}},")

        lines.append("}")
        return "\n".join(lines)

    # ── preview ───────────────────────────────────────────────────────────
    def _update_preview(self):
        entries = self._get_entries()[:8]
        abbrev = self._abbrev_check.isChecked()
        fmt_idx = self._format_combo.currentIndex()

        if fmt_idx == 0 or fmt_idx == 3:  # DOCX or PDF
            style = self._current_style()
            lines = []
            for e in entries:
                e_copy = dict(e)
                if abbrev:
                    e_copy["authors"] = self._abbreviate_authors(e_copy.get("authors", ""))
                lines.append(CitationStyleEngine.format_plain(e_copy, style))
            self._preview.setPlainText("\n\n".join(lines) if lines else "(no entries)")
        elif fmt_idx == 1:  # BibTeX
            keys = self._deduplicate_keys(entries)
            blocks = []
            for e, key in zip(entries, keys):
                blocks.append(self._format_entry_bibtex(e, key, abbrev))
            self._preview.setPlainText("\n\n".join(blocks) if blocks else "(no entries)")
        else:  # BibLaTeX
            keys = self._deduplicate_keys(entries)
            blocks = []
            for e, key in zip(entries, keys):
                blocks.append(self._format_entry_biblatex(e, key, abbrev))
            self._preview.setPlainText("\n\n".join(blocks) if blocks else "(no entries)")

    # ── export ────────────────────────────────────────────────────────────
    def _do_export(self):
        self._save_presets()
        fmt_idx = self._format_combo.currentIndex()
        if fmt_idx == 0:
            self._export_docx()
        elif fmt_idx == 1:
            self._export_bibtex()
        elif fmt_idx == 2:
            self._export_biblatex()
        else:
            self._export_pdf()

    def _export_docx(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save DOCX", str(Path.home() / "references.docx"),
            "Word Document (*.docx)"
        )
        if not path:
            return
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc        = Document()
            style      = doc.styles["Normal"]
            style.font.name = self._font_combo.currentText()
            style.font.size = Pt(self._size_spin.value())

            h = doc.add_heading(self._title_edit.text(), 0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER

            sub   = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r     = sub.add_run(f"Style: {self._current_style()}")
            r.font.size   = Pt(10)
            r.font.italic = True

            entries   = self._get_entries()
            cit_style = self._current_style()
            abbrev    = self._abbrev_check.isChecked()

            for e in entries:
                para = doc.add_paragraph()
                para.paragraph_format.space_after        = Pt(6)
                para.paragraph_format.first_line_indent  = Pt(-18)
                para.paragraph_format.left_indent        = Pt(18)

                e_copy = dict(e)
                if abbrev:
                    e_copy["authors"] = self._abbreviate_authors(e_copy.get("authors", ""))

                for text, bold, is_italic in CitationStyleEngine.format_runs(e_copy, cit_style):
                    run             = para.add_run(text)
                    run.font.name   = self._font_combo.currentText()
                    run.font.size   = Pt(self._size_spin.value())
                    run.font.italic = is_italic
                    run.font.bold   = bold

            doc.save(path)
            QMessageBox.information(self, self._lang.tr("docx_export_success_title"), 
                                    self._lang.tr("docx_export_success_text").format(path=path))
            self.accept()
        except Exception as ex:
            QMessageBox.critical(self, self._lang.tr("docx_export_fail_title"), str(ex))

    def _export_bibtex(self):
        path, _ = QFileDialog.getSaveFileName(
            self, self._lang.tr("bibtex_export_save_title"),
            str(Path.home() / "references.bib"),
            "BibTeX (*.bib)"
        )
        if not path:
            return
        try:
            entries = self._get_entries()
            abbrev = self._abbrev_check.isChecked()
            keys = self._deduplicate_keys(entries)
            blocks = []
            for e, key in zip(entries, keys):
                blocks.append(self._format_entry_bibtex(e, key, abbrev))

            content = "\n\n".join(blocks) + "\n"
            Path(path).write_text(content, encoding="utf-8")

            QMessageBox.information(
                self, self._lang.tr("bibtex_export_success_title"),
                self._lang.tr("bibtex_export_success_text").format(path=path)
            )
            self.accept()
        except Exception as ex:
            QMessageBox.critical(
                self, self._lang.tr("bibtex_export_fail_title"), str(ex)
            )

    def _export_biblatex(self):
        path, _ = QFileDialog.getSaveFileName(
            self, self._lang.tr("biblatex_export_save_title"),
            str(Path.home() / "references.bib"),
            "BibLaTeX (*.bib)"
        )
        if not path:
            return
        try:
            entries = self._get_entries()
            abbrev = self._abbrev_check.isChecked()
            keys = self._deduplicate_keys(entries)
            blocks = []
            for e, key in zip(entries, keys):
                blocks.append(self._format_entry_biblatex(e, key, abbrev))

            content = "\n\n".join(blocks) + "\n"
            Path(path).write_text(content, encoding="utf-8")

            QMessageBox.information(
                self, self._lang.tr("biblatex_export_success_title"),
                self._lang.tr("biblatex_export_success_text").format(path=path)
            )
            self.accept()
        except Exception as ex:
            QMessageBox.critical(
                self, self._lang.tr("biblatex_export_fail_title"), str(ex)
            )

    def _export_pdf(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save PDF", str(Path.home() / "references.pdf"),
            "PDF Document (*.pdf)"
        )
        if not path:
            return
            
        try:
            from PySide6.QtGui import QTextDocument, QFont
            from PySide6.QtPrintSupport import QPrinter
            
            doc = QTextDocument()
            font = QFont(self._font_combo.currentText(), self._size_spin.value())
            doc.setDefaultFont(font)
            
            html = f"<h1 align='center'>{self._title_edit.text()}</h1>"
            html += f"<p align='center'><i>Style: {self._current_style()}</i></p><br>"
            
            entries = self._get_entries()
            cit_style = self._current_style()
            abbrev = self._abbrev_check.isChecked()
            
            for e in entries:
                e_copy = dict(e)
                if abbrev:
                    e_copy["authors"] = self._abbreviate_authors(e_copy.get("authors", ""))
                    
                html += "<p style='margin-bottom: 12px;'>"
                for text, bold, is_italic in CitationStyleEngine.format_runs(e_copy, cit_style):
                    # Escape HTML for text
                    import html as html_lib
                    safe_text = html_lib.escape(text)
                    if bold:
                        safe_text = f"<b>{safe_text}</b>"
                    if is_italic:
                        safe_text = f"<i>{safe_text}</i>"
                    html += safe_text
                html += "</p>"
                
            doc.setHtml(html)
            
            printer = QPrinter(QPrinter.HighResolution)
            printer.setOutputFormat(QPrinter.PdfFormat)
            printer.setOutputFileName(path)
            
            doc.print_(printer)
            
            QMessageBox.information(
                self, "Export Successful",
                f"PDF saved successfully to {path}"
            )
            self.accept()
        except Exception as ex:
            QMessageBox.critical(self, "Export Failed", str(ex))
